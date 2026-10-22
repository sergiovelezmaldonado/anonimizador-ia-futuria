"""Fixtures y validación de descargas reales; solo datos ficticios."""
import json, sys, re, io, zipfile
from pathlib import Path
from docx import Document
from PIL import Image
from pypdf import PdfReader

if sys.argv[1] == 'fixture':
    out = Path(sys.argv[2]); out.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.author = 'AUTOR_CONFIDENCIAL_DE_PRUEBA'
    doc.core_properties.title = 'TITULO_PRIVADO_DE_PRUEBA'
    doc.add_heading('Informe de evaluación', 1)
    p = doc.add_paragraph('Cliente: '); p.add_run('Ana Ruiz').bold = True
    p = doc.add_paragraph('Correo: '); p.add_run('ana@ejemplo.test').italic = True
    doc.add_paragraph('')
    p = doc.add_paragraph('Condición: '); p.add_run('revisión ágil').bold = True; p.add_run(' y envío mañana.').italic = True
    doc.add_paragraph('Primera tarea', style='List Bullet')
    doc.add_paragraph('Segunda tarea', style='List Bullet')
    image = io.BytesIO(); Image.new('RGB', (12, 12), '#b00020').save(image, format='PNG'); image.seek(0)
    doc.add_picture(image)
    doc.add_paragraph('FIN_DOCUMENTO_XZ')
    doc.save(str(out))
    print(json.dumps({'fixture': str(out), 'bytes': out.stat().st_size}))
else:
    report_path = Path(sys.argv[2]); report = json.loads(report_path.read_text(encoding='utf-8'))
    result = []
    for item in report['artifacts']:
        p = Path(item['path']); expected = item['text']; fmt = item['format']
        assert p.exists() and p.stat().st_size > 0, str(p)
        info = {'file': p.name, 'format': fmt, 'bytes': p.stat().st_size}
        if fmt == 'txt':
            with p.open(encoding='utf-8', newline='') as f: actual = f.read()
            assert actual == expected, 'TXT inexacto'
        elif fmt == 'json':
            data = json.loads(p.read_text(encoding='utf-8'))
            assert set(data) == {'format', 'version', 'kind', 'document'}
            assert data['format'] == 'futuria-text' and data['version'] == 1
            assert data['document']['text'] == expected and data['kind'] == item['kind']
            actual = p.read_text(encoding='utf-8')
        elif fmt == 'docx':
            doc = Document(str(p))
            actual = '\n'.join(para.text for para in doc.paragraphs)
            assert actual == expected.replace('\r\n', '\n').replace('\r', '\n'), 'DOCX omite/repite texto'
            if item.get('rich'):
                heading_style = doc.paragraphs[0].style
                assert heading_style is not None and heading_style.name == 'Heading 1', 'sin título'
                assert any(r.bold for para in doc.paragraphs for r in para.runs), 'sin negrita'
                assert any(r.italic for para in doc.paragraphs for r in para.runs), 'sin cursiva'
            with zipfile.ZipFile(p) as z:
                payload = '\n'.join(z.read(n).decode('utf-8', errors='replace') for n in z.namelist())
                assert not any(n.startswith('word/media/') for n in z.namelist()), 'imágenes en salida textual'
                for value in item.get('forbidden', []): assert value not in payload, 'dato privado en DOCX'
        elif fmt == 'pdf':
            reader = PdfReader(p); actual = '\n'.join(page.extract_text() or '' for page in reader.pages)
            assert re.sub(r'\s+', '', actual) == re.sub(r'\s+', '', expected), 'PDF omite/repite texto'
            info['pages'] = len(reader.pages)
        else: raise AssertionError(fmt)
        if fmt != 'docx':
            for value in item.get('forbidden', []): assert value not in actual, 'dato privado en salida limpia'
        if item.get('markers'):
            markers = re.findall(r'MARCADOR_PAGINA_\d{3}', actual)
            assert markers == [f'MARCADOR_PAGINA_{i:03d}' for i in range(1, item['markers'] + 1)], 'marcadores incompletos'
            info['markers'] = len(markers)
        result.append(info)
    report['artifactVerification'] = {'passed': len(result), 'artifacts': result}
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report['artifactVerification'], ensure_ascii=False, indent=2))
